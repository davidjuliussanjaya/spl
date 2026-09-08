from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor
from PIL import Image


ROOT = Path(__file__).resolve().parents[1]
IMAGE_DIR = ROOT / "docs" / "SC"
OUTPUT = ROOT / "docs" / "User_Manual_Admin_SPL_David_Julius_Sanjaya_22410100009.docx"
REFERENCE_COVER = Path.home() / "AppData" / "Local" / "Temp" / "buku_manual_page_1.png"
LOGO_FILE = Path.home() / "AppData" / "Local" / "Temp" / "manual_admin_undika_logo.png"

MAROON = "991B2D"
LIGHT_MAROON = "FCECEF"
GOLD = "C9A227"
NAVY = "16366F"
GRAY = "666666"
LIGHT_GRAY = "F2F2F2"


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=100, bottom=100, end=100):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_page_numbering(section, start, fmt="decimal"):
    sect_pr = section._sectPr
    pg_num = sect_pr.find(qn("w:pgNumType"))
    if pg_num is None:
        pg_num = OxmlElement("w:pgNumType")
        sect_pr.append(pg_num)
    pg_num.set(qn("w:start"), str(start))
    pg_num.set(qn("w:fmt"), fmt)


def add_field(paragraph, instruction, placeholder=""):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = placeholder
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend((begin, instr, separate, text, end))
    return run


def add_page_number(section):
    section.header.is_linked_to_previous = False
    paragraph = section.header.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    paragraph.paragraph_format.space_after = Pt(0)
    add_field(paragraph, " PAGE ", "1")


def add_body(document, text, *, bold_prefix=None, first_line=True, align=WD_ALIGN_PARAGRAPH.JUSTIFY):
    paragraph = document.add_paragraph()
    paragraph.alignment = align
    paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    paragraph.paragraph_format.space_after = Pt(6)
    if first_line:
        paragraph.paragraph_format.first_line_indent = Cm(1.25)
    if bold_prefix and text.startswith(bold_prefix):
        paragraph.add_run(bold_prefix).bold = True
        paragraph.add_run(text[len(bold_prefix):])
    else:
        paragraph.add_run(text)
    return paragraph


def add_steps(document, title, steps):
    heading = document.add_paragraph()
    heading.paragraph_format.keep_with_next = True
    heading.paragraph_format.space_before = Pt(5)
    heading.paragraph_format.space_after = Pt(3)
    heading.add_run(title).bold = True
    for number, step in enumerate(steps, 1):
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        paragraph.paragraph_format.left_indent = Cm(0.65)
        paragraph.paragraph_format.first_line_indent = Cm(-0.65)
        paragraph.paragraph_format.line_spacing = 1.15
        paragraph.paragraph_format.space_after = Pt(3)
        paragraph.add_run(f"{number}. ").bold = True
        paragraph.add_run(step)


def add_note(document, text, label="Catatan"):
    table = document.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    cell = table.cell(0, 0)
    shade_cell(cell, LIGHT_MAROON)
    set_cell_margins(cell, top=120, start=150, bottom=120, end=150)
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.1
    run = paragraph.add_run(f"{label}: ")
    run.bold = True
    run.font.color.rgb = RGBColor.from_string(MAROON)
    paragraph.add_run(text)
    document.add_paragraph().paragraph_format.space_after = Pt(0)


def add_figure(document, image_name, caption, explanation, width=Inches(6.35)):
    path = IMAGE_DIR / image_name
    if not path.exists():
        raise FileNotFoundError(path)
    picture_paragraph = document.add_paragraph()
    picture_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    picture_paragraph.paragraph_format.space_before = Pt(4)
    picture_paragraph.paragraph_format.space_after = Pt(2)
    picture_paragraph.paragraph_format.keep_with_next = True
    picture_paragraph.add_run().add_picture(str(path), width=width)

    caption_paragraph = document.add_paragraph()
    caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption_paragraph.paragraph_format.space_before = Pt(0)
    caption_paragraph.paragraph_format.space_after = Pt(6)
    caption_paragraph.paragraph_format.keep_with_next = True
    run = caption_paragraph.add_run(caption)
    run.font.name = "Times New Roman"
    run.font.size = Pt(11)

    add_body(document, explanation, first_line=True)


def add_data_table(document, headers, rows, widths=None):
    table = document.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.autofit = True
    header = table.rows[0]
    set_repeat_table_header(header)
    for i, value in enumerate(headers):
        cell = header.cells[i]
        shade_cell(cell, MAROON)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margins(cell)
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run(value)
        run.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(values):
            if row_index % 2:
                shade_cell(cells[i], LIGHT_GRAY)
            set_cell_margins(cells[i])
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            paragraph = cells[i].paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.add_run(str(value))
    if widths:
        for row in table.rows:
            for i, width in enumerate(widths):
                row.cells[i].width = width
    document.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def prepare_logo():
    if not REFERENCE_COVER.exists():
        return None
    with Image.open(REFERENCE_COVER) as source:
        logo = source.crop((165, 135, 545, 255))
        logo.save(LOGO_FILE)
    return LOGO_FILE


def configure_styles(document):
    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.font.size = Pt(12)
    normal.paragraph_format.space_after = Pt(6)

    for style_name, size, color, before, after in (
        ("Title", 18, "000000", 0, 10),
        ("Heading 1", 14, "000000", 10, 8),
        ("Heading 2", 12, "000000", 8, 5),
        ("Heading 3", 12, MAROON, 6, 4),
    ):
        style = document.styles[style_name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
    document.styles["Heading 1"].paragraph_format.page_break_before = True

    if "Manual Caption" not in document.styles:
        caption_style = document.styles.add_style("Manual Caption", WD_STYLE_TYPE.PARAGRAPH)
        caption_style.font.name = "Times New Roman"
        caption_style.font.size = Pt(11)


def configure_section(section):
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(3)
    section.right_margin = Cm(2.5)
    section.header_distance = Cm(1.2)
    section.footer_distance = Cm(1.2)


def add_cover(document):
    section = document.sections[0]
    configure_section(section)
    section.different_first_page_header_footer = True

    logo = prepare_logo()
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.space_after = Pt(14)
    if logo:
        paragraph.add_run().add_picture(str(logo), width=Inches(3.2))

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title.paragraph_format.space_after = Pt(0)
    for line in ("APLIKASI SURVEY PENGGUNA LULUSAN", "UNIVERSITAS DINAMIKA"):
        run = title.add_run(line)
        run.bold = True
        run.font.name = "Times New Roman"
        run.font.size = Pt(13)
        run.add_break()

    manual = document.add_paragraph()
    manual.paragraph_format.space_before = Pt(42)
    manual.paragraph_format.space_after = Pt(62)
    run = manual.add_run("MANUAL BOOK")
    run.bold = True
    run.font.size = Pt(13)

    program = document.add_paragraph()
    program.paragraph_format.space_after = Pt(72)
    for index, line in enumerate(("Program Studi", "S1 Sistem Informasi")):
        run = program.add_run(line)
        run.bold = True
        run.font.size = Pt(12)
        if index == 0:
            run.add_break()

    author = document.add_paragraph()
    author.paragraph_format.space_after = Pt(60)
    for index, line in enumerate(("Oleh:", "DAVID JULIUS SANJAYA", "22410100009")):
        run = author.add_run(line)
        run.bold = True
        run.font.size = Pt(12)
        if index < 2:
            run.add_break()

    line = document.add_paragraph()
    line.paragraph_format.space_before = Pt(12)
    line.paragraph_format.space_after = Pt(4)
    p_pr = line._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    top = OxmlElement("w:top")
    top.set(qn("w:val"), "double")
    top.set(qn("w:sz"), "12")
    top.set(qn("w:space"), "2")
    top.set(qn("w:color"), "FF0000")
    borders.append(top)
    p_pr.append(borders)

    institution = document.add_paragraph()
    institution.paragraph_format.space_after = Pt(0)
    lines = ("FAKULTAS TEKNOLOGI DAN INFORMATIKA", "UNIVERSITAS DINAMIKA", "2026")
    for index, text in enumerate(lines):
        run = institution.add_run(text)
        run.bold = True
        run.font.size = Pt(12)
        if index < len(lines) - 1:
            run.add_break()


def add_toc(document):
    section = document.add_section(WD_SECTION_START.NEW_PAGE)
    configure_section(section)
    set_page_numbering(section, 2, "lowerRoman")
    add_page_number(section)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(18)
    run = title.add_run("DAFTAR ISI")
    run.bold = True
    run.font.size = Pt(14)

    toc = document.add_paragraph()
    toc.paragraph_format.space_after = Pt(0)
    add_field(toc, ' TOC \\o "1-3" \\h \\z \\u ', "Daftar isi akan diperbarui saat dokumen dibuka.")


def add_intro(document):
    section = document.add_section(WD_SECTION_START.NEW_PAGE)
    configure_section(section)
    set_page_numbering(section, 1, "decimal")
    add_page_number(section)

    document.add_heading("BAB I\nPENDAHULUAN", level=1).alignment = WD_ALIGN_PARAGRAPH.CENTER
    document.add_heading("1.1 Tentang Aplikasi SPL", level=2)
    add_body(
        document,
        "SPL (Survey Pengguna Lulusan) adalah aplikasi Tracer Study Universitas Dinamika yang digunakan untuk mengumpulkan, mengelola, dan menganalisis penilaian perusahaan atau instansi terhadap kinerja lulusan. Sistem menyediakan kode akses survey bagi penyelia tanpa mewajibkan pembuatan akun responden.",
    )
    document.add_heading("1.2 Tujuan Manual Book", level=2)
    add_body(
        document,
        "Manual book ini menjadi panduan bagi admin dalam mengelola data master, menyusun survey, memantau hasil evaluasi, melihat arsip, serta mengunduh laporan. Seluruh tampilan dan prosedur dalam dokumen ini berasal dari proyek SPL.",
    )
    document.add_heading("1.3 Hak Akses Admin", level=2)
    add_data_table(
        document,
        ("Menu", "Fungsi Utama"),
        (
            ("Dashboard", "Melihat ringkasan dan analisis hasil evaluasi lulusan."),
            ("Survey", "Membuat survey tunggal atau massal dan memantau status pengisian."),
            ("Lulusan", "Mengelola data lulusan yang akan dinilai."),
            ("Pengguna Lulusan", "Mengelola perusahaan, instansi, dan data penyelia."),
            ("Pertanyaan", "Mengelola pertanyaan serta pilihan jawaban survey."),
            ("Aspek Evaluasi", "Mengelola kategori atau aspek penilaian."),
            ("Cetak Laporan", "Mengunduh rekap hasil survey dalam format Excel."),
            ("Arsip Survey", "Melihat survey yang telah selesai dan tersimpan permanen."),
        ),
        widths=(Cm(4), Cm(11)),
    )
    document.add_heading("1.4 Alur Kerja Admin", level=2)
    add_steps(
        document,
        "Urutan penggunaan yang disarankan",
        (
            "Tambahkan data perusahaan atau instansi pada menu Pengguna Lulusan.",
            "Tambahkan data lulusan dan hubungkan dengan perusahaan tempat lulusan bekerja.",
            "Siapkan aspek evaluasi dan pertanyaan survey yang akan digunakan.",
            "Buat survey tunggal atau survey massal, kemudian bagikan kode akses kepada penyelia.",
            "Pantau status survey, analisis hasil pada Dashboard, dan periksa Arsip Survey.",
            "Unduh laporan Excel secara berkala untuk kebutuhan dokumentasi.",
        ),
    )


def add_access_and_dashboard(document):
    document.add_heading("BAB II\nAKSES SISTEM DAN DASHBOARD", level=1).alignment = WD_ALIGN_PARAGRAPH.CENTER
    document.add_heading("2.1 Login Admin", level=2)
    add_figure(
        document,
        "Screenshot 2026-08-31 115921.png",
        "Gambar L3.1 Halaman Login Admin",
        "Gambar L3.1 menampilkan halaman autentikasi SPL. Admin menggunakan tab Login Admin, kemudian memasukkan email dan password sesuai akun yang telah terdaftar sebelum masuk ke Dashboard.",
    )
    add_steps(
        document,
        "Langkah penggunaan",
        (
            "Buka alamat aplikasi SPL melalui peramban.",
            "Pilih tab Login Admin.",
            "Masukkan email staff dan password.",
            "Aktifkan Ingat Saya hanya pada perangkat pribadi.",
            "Klik Masuk ke Dashboard dan tunggu proses validasi selesai.",
        ),
    )
    add_note(document, "Jangan menuliskan password pada dokumen atau membagikannya kepada pihak lain.")

    document.add_heading("2.2 Dashboard Admin", level=2)
    add_figure(
        document,
        "Screenshot 2026-08-31 115942.png",
        "Gambar L3.2 Halaman Dashboard Evaluasi Lulusan",
        "Gambar L3.2 menunjukkan Dashboard admin yang merangkum jumlah lulusan yang telah dinilai, indeks kepuasan, kategori terbaik dan terendah, sebaran responden berdasarkan program studi, rata-rata penilaian per kategori, serta umpan balik terbaru.",
    )
    add_steps(
        document,
        "Cara membaca Dashboard",
        (
            "Periksa kartu Lulusan Dinilai untuk mengetahui jumlah respon survey yang telah diarsipkan.",
            "Gunakan Indeks Kepuasan untuk melihat rata-rata skor keseluruhan pada skala 4,00.",
            "Tinjau Kategori Terbaik dan Kategori Terendah sebagai dasar evaluasi.",
            "Baca grafik responden berdasarkan program studi dan grafik rata-rata kategori.",
            "Gunakan tombol Lihat Selengkapnya atau Lihat Semua untuk membuka informasi yang lebih rinci.",
        ),
    )

    document.add_heading("2.3 Memfilter Data Dashboard", level=2)
    add_figure(
        document,
        "Screenshot 2026-08-31 115957.png",
        "Gambar L3.3 Hasil Filter Dashboard",
        "Gambar L3.3 memperlihatkan Dashboard setelah filter periode 2026, Fakultas Desain dan Industri Kreatif, dan Program Studi Desain Komunikasi Visual diterapkan. Statistik dan grafik otomatis menyesuaikan data yang dipilih.",
    )
    add_steps(
        document,
        "Langkah penggunaan filter",
        (
            "Pilih satu atau beberapa periode pada filter Periode.",
            "Pilih Fakultas untuk mempersempit data.",
            "Pilih Program Studi yang tersedia pada fakultas tersebut.",
            "Klik Terapkan.",
            "Periksa label filter aktif dan pastikan statistik telah berubah sesuai pilihan.",
            "Gunakan tombol reset untuk kembali menampilkan seluruh data.",
        ),
    )
    add_note(document, "Dashboard menggunakan data Arsip Survey. Survey yang belum selesai tidak masuk ke perhitungan.")


def add_survey(document):
    document.add_heading("BAB III\nMANAJEMEN SURVEY", level=1).alignment = WD_ALIGN_PARAGRAPH.CENTER
    document.add_heading("3.1 Daftar dan Status Survey", level=2)
    add_figure(
        document,
        "Screenshot 2026-08-31 120010.png",
        "Gambar L3.4 Halaman Manajemen Survey",
        "Gambar L3.4 menampilkan daftar sesi survey beserta judul, tahun, perusahaan, lulusan terkait, kode akses, status, dan tombol aksi. Status Belum Diisi berarti survey masih menunggu responden, sedangkan Selesai berarti jawaban telah dikirim dan disimpan ke arsip.",
    )
    add_steps(
        document,
        "Operasi pada daftar survey",
        (
            "Gunakan kolom Judul Survey, Nama Perusahaan, Alumni/Lulusan, atau Status untuk mencari data.",
            "Salin kode akses dan kirimkan kepada penyelia yang berhak mengisi survey.",
            "Klik ikon mata untuk melihat detail survey.",
            "Klik ikon tempat sampah hanya jika survey benar-benar tidak diperlukan.",
            "Gunakan Buat Survey Baru untuk satu lulusan atau Buat Survey Massal untuk beberapa lulusan.",
        ),
    )
    add_note(document, "Kode akses bersifat unik. Pastikan kode diberikan kepada penyelia dan perusahaan yang sesuai.")

    document.add_heading("3.2 Membuat Survey Tunggal", level=2)
    add_figure(
        document,
        "Screenshot 2026-08-31 120021.png",
        "Gambar L3.5 Halaman Pembuatan Survey Tunggal",
        "Gambar L3.5 menunjukkan formulir pembuatan survey untuk satu lulusan. Form mencakup informasi utama survey, identitas alumni, data penyelia, identitas perusahaan, dan pilihan pertanyaan yang akan digunakan.",
    )
    add_steps(
        document,
        "Langkah pembuatan",
        (
            "Klik Buat Survey Baru pada halaman Manajemen Survey.",
            "Isi Judul Survey, Tahun Survey, dan deskripsi atau instruksi bila diperlukan.",
            "Pilih data lulusan yang akan dinilai.",
            "Pilih instansi atau perusahaan. Data penyelia dan perusahaan akan terisi dari data master bila tersedia.",
            "Periksa kembali nomor HP, email, nomor badan hukum, telepon, dan alamat perusahaan.",
            "Pilih minimal satu pertanyaan aktif yang sesuai dengan fakultas lulusan.",
            "Simpan survey untuk menghasilkan kode akses.",
        ),
    )

    document.add_heading("3.3 Membuat Survey Massal", level=2)
    add_figure(
        document,
        "Screenshot 2026-08-31 120040.png",
        "Gambar L3.6 Halaman Pembuatan Survey Massal",
        "Gambar L3.6 menampilkan formulir survey massal yang digunakan untuk membuat survey sekaligus bagi seluruh lulusan pada tahun lulus tertentu. Sistem melakukan preview sebelum pembuatan agar admin dapat memeriksa kelengkapan lulusan dan perusahaan.",
    )
    add_steps(
        document,
        "Langkah pembuatan",
        (
            "Klik Buat Survey Massal pada halaman Manajemen Survey.",
            "Isi judul, tahun survey, dan deskripsi atau instruksi.",
            "Pilih Tahun Lulus.",
            "Klik Lihat Preview Lulusan.",
            "Periksa daftar lulusan yang ditemukan dan pastikan setiap lulusan memiliki perusahaan terkait.",
            "Pilih pertanyaan yang akan digunakan.",
            "Klik tombol pembuatan survey untuk seluruh lulusan yang valid.",
        ),
    )
    add_note(document, "Lulusan yang belum terhubung dengan perusahaan tidak dapat dibuatkan survey massal.")


def add_master_data(document):
    document.add_heading("BAB IV\nPENGELOLAAN DATA MASTER", level=1).alignment = WD_ALIGN_PARAGRAPH.CENTER
    document.add_heading("4.1 Daftar Lulusan", level=2)
    add_figure(
        document,
        "Screenshot 2026-08-31 120059.png",
        "Gambar L3.7 Halaman Data Lulusan",
        "Gambar L3.7 menampilkan daftar lulusan Universitas Dinamika. Admin dapat mencari berdasarkan nama, NIM, fakultas, program studi, rentang tahun lulus, dan status, serta mengekspor data ke Excel.",
    )
    add_steps(
        document,
        "Penggunaan halaman",
        (
            "Masukkan parameter pencarian atau pilih filter yang diperlukan.",
            "Klik Cari untuk menerapkan filter atau Reset untuk menghapusnya.",
            "Klik ikon mata pada kolom Aksi untuk melihat detail lulusan dan perusahaan terkait.",
            "Klik Export Excel untuk mengunduh daftar lulusan.",
            "Klik Tambah Lulusan untuk membuat data baru.",
        ),
    )

    document.add_heading("4.2 Menambah Data Lulusan", level=2)
    add_figure(
        document,
        "Screenshot 2026-08-31 120113.png",
        "Gambar L3.8 Halaman Tambah Data Lulusan",
        "Gambar L3.8 menunjukkan formulir penambahan lulusan yang terdiri dari informasi akademik dan pengaturan data. Lulusan perlu dihubungkan dengan perusahaan atau pengguna lulusan agar dapat dipilih pada survey.",
    )
    add_steps(
        document,
        "Langkah penambahan",
        (
            "Isi Nama Lengkap dan NIM yang belum pernah digunakan.",
            "Pilih Fakultas dan Program Studi.",
            "Masukkan tanggal atau tahun lulus.",
            "Pilih Perusahaan/Pengguna Lulusan yang menilai lulusan tersebut.",
            "Aktifkan Status Lulusan bila data masih digunakan.",
            "Klik Simpan Data.",
        ),
    )

    document.add_heading("4.3 Daftar Pengguna Lulusan", level=2)
    add_figure(
        document,
        "Screenshot 2026-08-31 120122.png",
        "Gambar L3.9 Halaman Pengguna Lulusan",
        "Gambar L3.9 menampilkan daftar perusahaan atau instansi penyerap lulusan. Informasi yang disajikan meliputi perusahaan, penyelia, kontak, jenis perusahaan, jumlah lulusan, cakupan wilayah, dan aksi pengelolaan.",
    )
    add_steps(
        document,
        "Operasi pada data perusahaan",
        (
            "Klik Tambah Pengguna untuk membuat data perusahaan baru.",
            "Klik ikon pensil untuk mengubah informasi perusahaan dan penyelia.",
            "Klik ikon tempat sampah untuk menghapus data setelah memastikan data tidak lagi digunakan.",
            "Periksa kolom Jumlah Lulusan untuk mengetahui keterkaitan perusahaan dengan data lulusan.",
        ),
    )

    document.add_heading("4.4 Menambah Pengguna Lulusan", level=2)
    add_figure(
        document,
        "Screenshot 2026-08-31 120134.png",
        "Gambar L3.10 Halaman Tambah Instansi atau Perusahaan",
        "Gambar L3.10 memperlihatkan formulir penambahan instansi yang mencakup informasi perusahaan, kontak penyelia atau atasan langsung, serta cakupan wilayah dan data pendukung.",
    )
    add_steps(
        document,
        "Langkah penambahan",
        (
            "Isi nama, jenis, nomor badan hukum, dan alamat perusahaan.",
            "Isi nama penyelia, jabatan, email, dan nomor WhatsApp atau HP.",
            "Masukkan jumlah cabang nasional dan luar negeri.",
            "Isi jumlah lulusan yang pernah bekerja serta durasi rata-rata bekerja bila datanya tersedia.",
            "Klik Simpan Instansi.",
        ),
    )
    add_note(document, "Email penyelia harus valid dan tidak boleh sama dengan data penyelia lain yang sudah tersimpan.")

    document.add_heading("4.5 Daftar Pertanyaan Survey", level=2)
    add_figure(
        document,
        "Screenshot 2026-08-31 120143.png",
        "Gambar L3.11 Halaman Data Pertanyaan",
        "Gambar L3.11 menampilkan pertanyaan survey beserta kategori, peruntukan fakultas, kode, tipe masukan, status wajib, dan status aktif. Hanya pertanyaan aktif yang dapat dipilih saat membuat survey baru.",
    )
    add_steps(
        document,
        "Pengelolaan pertanyaan",
        (
            "Gunakan kolom pencarian, filter tipe soal, dan filter status untuk menemukan pertanyaan.",
            "Klik Buat Pertanyaan untuk menambahkan pertanyaan baru.",
            "Klik ikon pensil untuk mengubah pertanyaan.",
            "Klik Matikan untuk menonaktifkan pertanyaan atau Hidupkan untuk mengaktifkannya kembali.",
        ),
    )

    document.add_heading("4.6 Menambah Pertanyaan", level=2)
    add_figure(
        document,
        "Screenshot 2026-08-31 120149.png",
        "Gambar L3.12 Halaman Tambah Pertanyaan",
        "Gambar L3.12 menunjukkan formulir pertanyaan baru. Admin menentukan teks pertanyaan, kategori, fakultas tujuan, tipe masukan, kode pertanyaan, status wajib, serta opsi dan nilai jawaban untuk tipe rating.",
    )
    add_steps(
        document,
        "Langkah penambahan",
        (
            "Tuliskan pertanyaan secara jelas dan tidak ambigu.",
            "Pilih kategori pertanyaan dan peruntukan fakultas. Pilih Umum bila berlaku untuk semua fakultas.",
            "Pilih tipe Pilihan Ganda/Rating atau Teks Bebas/Essay.",
            "Isi kode pertanyaan bila pertanyaan mengikuti kode standar.",
            "Aktifkan Required bila responden wajib menjawab.",
            "Untuk tipe rating, tambahkan opsi jawaban dan nilai numeriknya.",
            "Simpan pertanyaan.",
        ),
    )
    add_note(document, "Nilai opsi rating digunakan pada perhitungan Dashboard dan laporan.")

    document.add_heading("4.7 Daftar Aspek Evaluasi", level=2)
    add_figure(
        document,
        "Screenshot 2026-08-31 120156.png",
        "Gambar L3.13 Halaman Data Aspek Evaluasi",
        "Gambar L3.13 menampilkan kategori yang digunakan untuk mengelompokkan pertanyaan, seperti Etika, Keahlian Berdasarkan Bidang Ilmu, Kemampuan Berbahasa Asing, Penggunaan Teknologi Informasi, dan Kepemimpinan.",
    )
    add_steps(
        document,
        "Pengelolaan aspek",
        (
            "Klik Tambah Kategori untuk membuat aspek evaluasi baru.",
            "Klik ikon pensil untuk mengubah nama atau deskripsi.",
            "Klik ikon tempat sampah untuk menghapus kategori yang tidak digunakan.",
            "Periksa keterkaitan kategori dengan pertanyaan sebelum melakukan penghapusan.",
        ),
    )

    document.add_heading("4.8 Menambah Aspek Evaluasi", level=2)
    add_figure(
        document,
        "Screenshot 2026-08-31 120204.png",
        "Gambar L3.14 Halaman Tambah Aspek Evaluasi",
        "Gambar L3.14 memperlihatkan formulir sederhana untuk menambahkan nama kategori dan deskripsi. Nama kategori menjadi label pengelompokan pada pertanyaan, Dashboard, arsip, dan laporan.",
    )
    add_steps(
        document,
        "Langkah penambahan",
        (
            "Klik Tambah Kategori dari halaman Data Kategori.",
            "Isi Nama Kategori.",
            "Tambahkan deskripsi singkat agar tujuan aspek mudah dipahami.",
            "Klik Simpan.",
        ),
    )


def add_reporting(document):
    document.add_heading("BAB V\nLAPORAN DAN ARSIP", level=1).alignment = WD_ALIGN_PARAGRAPH.CENTER
    document.add_heading("5.1 Mengunduh Laporan Excel", level=2)
    add_figure(
        document,
        "Screenshot 2026-08-31 120211.png",
        "Gambar L3.15 Halaman Cetak Laporan Tracer Study",
        "Gambar L3.15 menampilkan filter laporan berdasarkan tahun lulus, fakultas, dan program studi. Halaman juga menjelaskan struktur file Excel berupa sheet per tahun, data per fakultas, dan tabel ringkasan distribusi.",
    )
    add_steps(
        document,
        "Langkah mengunduh laporan",
        (
            "Buka menu Cetak Laporan.",
            "Pilih Tahun Lulus atau biarkan Semua Tahun untuk membuat satu sheet per tahun.",
            "Pilih Fakultas dan Program Studi bila laporan perlu dipersempit.",
            "Klik Download Excel.",
            "Buka file hasil unduhan dan periksa sheet serta kolom yang tersedia.",
        ),
    )
    add_note(document, "Laporan hanya memuat survey yang sudah selesai. Gunakan Reset Filter bila hasil unduhan kosong atau terlalu spesifik.")

    document.add_heading("5.2 Melihat Arsip Survey", level=2)
    add_figure(
        document,
        "Screenshot 2026-08-31 120221.png",
        "Gambar L3.16 Halaman Arsip Survey",
        "Gambar L3.16 menunjukkan daftar survey yang telah selesai dan tersimpan permanen sebagai arsip evaluasi lulusan. Admin dapat memfilter arsip berdasarkan kata kunci, tahun instrumen, fakultas, dan program studi.",
    )
    add_steps(
        document,
        "Langkah penggunaan",
        (
            "Masukkan nama, NIM, perusahaan, atau penyelia pada kolom Pencarian bila diperlukan.",
            "Pilih Tahun Instrumen, Fakultas, dan Program Studi.",
            "Klik Terapkan atau Reset untuk menghapus filter.",
            "Klik Lihat pada baris arsip untuk membuka identitas dan jawaban survey.",
            "Pada halaman detail, gunakan Cetak/Simpan PDF bila diperlukan.",
        ),
    )
    add_note(document, "Arsip merupakan snapshot saat survey disubmit. Perubahan data master setelahnya tidak mengubah isi arsip.")

    document.add_heading("5.3 Memeriksa Hasil Laporan", level=2)
    add_figure(
        document,
        "Screenshot 2026-08-31 120236.png",
        "Gambar L3.17 Hasil Laporan Tracer Study dalam Excel",
        "Gambar L3.17 memperlihatkan hasil ekspor Excel yang membagi data berdasarkan fakultas dan tahun lulus. Di bagian bawah setiap sheet tersedia ringkasan distribusi penilaian per program studi dan kode soal.",
    )
    add_steps(
        document,
        "Pemeriksaan hasil",
        (
            "Pastikan nama file dan tahun laporan sesuai filter yang dipilih.",
            "Buka setiap sheet tahun lulus yang tersedia.",
            "Periksa bagian fakultas FTI, FDIK, dan FEB sesuai data.",
            "Pastikan identitas alumni, responden, perusahaan, dan nilai per kode soal tampil.",
            "Tinjau tabel ringkasan distribusi untuk melihat persentase Sangat Baik, Baik, Kurang, dan Sangat Kurang.",
        ),
    )


def add_troubleshooting(document):
    document.add_heading("BAB VI\nPENANGANAN MASALAH", level=1).alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_body(
        document,
        "Bagian ini merangkum masalah operasional yang umum ditemui oleh admin dan tindakan pemeriksaan awal yang dapat dilakukan sebelum meminta bantuan teknis.",
    )
    add_data_table(
        document,
        ("Masalah", "Penyebab Umum", "Solusi"),
        (
            ("Menu admin tidak tampil", "Akun tidak memiliki role admin.", "Gunakan akun admin atau periksa role akun pada basis data."),
            ("Login gagal", "Email/password salah atau akun tidak aktif.", "Periksa ejaan, Caps Lock, dan status akun."),
            ("Survey massal kosong", "Tidak ada lulusan sesuai tahun atau perusahaan belum terhubung.", "Lengkapi data lulusan dan pengguna lulusan."),
            ("Pertanyaan tidak tersedia", "Pertanyaan nonaktif atau fakultas tidak sesuai.", "Aktifkan pertanyaan dan periksa peruntukan fakultas."),
            ("Kode akses ditolak", "Kode salah atau survey sudah selesai.", "Periksa kode dan status pada halaman Survey."),
            ("Dashboard/laporan kosong", "Belum ada survey selesai atau filter terlalu sempit.", "Periksa Arsip Survey dan reset filter."),
            ("Survey tidak bisa diubah", "Survey sudah selesai dan dikunci.", "Gunakan detail survey atau arsip untuk melihat hasil."),
        ),
        widths=(Cm(4), Cm(5), Cm(7)),
    )

    document.add_heading("6.1 Rekomendasi Operasional", level=2)
    add_steps(
        document,
        "Praktik yang disarankan",
        (
            "Masukkan perusahaan sebelum data lulusan.",
            "Pastikan setiap lulusan yang akan disurvey terhubung dengan perusahaan.",
            "Gunakan pertanyaan Umum untuk aspek lintas fakultas.",
            "Periksa preview sebelum membuat survey massal.",
            "Simpan catatan kode akses yang telah dibagikan.",
            "Unduh laporan secara berkala sebagai dokumentasi administratif.",
            "Gunakan Arsip Survey sebagai sumber resmi untuk survey yang telah selesai.",
        ),
    )


def build_document():
    images = sorted(IMAGE_DIR.glob("*.png"))
    if len(images) != 17:
        raise RuntimeError(f"Expected 17 PNG screenshots in {IMAGE_DIR}, found {len(images)}")

    document = Document()
    document.core_properties.title = "User Manual Admin SPL"
    document.core_properties.subject = "Panduan penggunaan SPL untuk role admin"
    document.core_properties.author = "David Julius Sanjaya"
    document.core_properties.keywords = "SPL, Survey Pengguna Lulusan, Tracer Study, Admin"
    configure_styles(document)

    add_cover(document)
    add_toc(document)
    add_intro(document)
    add_access_and_dashboard(document)
    add_survey(document)
    add_master_data(document)
    add_reporting(document)
    add_troubleshooting(document)

    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_document()
